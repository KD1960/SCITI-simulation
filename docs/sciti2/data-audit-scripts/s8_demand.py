import openpyxl, pandas as pd, numpy as np, re, collections
pd.set_option('display.width',250); pd.set_option('display.max_columns',60); pd.set_option('display.max_rows',100)
X="/Users/kevindooley/Docs/School/LE/SCM and AI Oct2026 workshop/Master Data Set_4.xlsx"
wb=openpyxl.load_workbook(X,read_only=True)
ws=wb["Demand_Retailer"]
rows=list(ws.iter_rows(min_row=1,max_row=268))
# formula patterns per column (rows 5-264), normalize row numbers
from openpyxl.utils import get_column_letter
pat=collections.defaultdict(collections.Counter)
for r in rows[4:264]:
    for c in r[:27]:
        v=c.value
        if isinstance(v,str) and v.startswith('='):
            pat[c.column_letter][re.sub(r'(?<=[A-Z])\d+',lambda m:'n' if m.group(0)!='2' else '2',v)]+=1
        else:
            pat[c.column_letter][type(v).__name__]+=1
for k in sorted(pat,key=lambda s:(len(s),s)): print(k,dict(pat[k]))
print("rows 265-268 and summary area AD:BB rows 1-30:")
for r in rows[264:268]:
    print([(c.coordinate,c.value) for c in r if c.value is not None])
for r in rows[0:30]:
    v=[(c.coordinate,c.value) for c in r[29:] if c.value is not None]
    if v: print(v)
# any RAND in any sheet?
for name in wb.sheetnames:
    cnt=0; ex=None
    for r in wb[name].iter_rows(max_row=600):
        for c in r:
            if isinstance(c.value,str) and 'RAND' in c.value.upper():
                cnt+=1; ex=ex or (c.coordinate,c.value)
    print("RAND formulas in first 600 rows of",name,cnt,ex)
