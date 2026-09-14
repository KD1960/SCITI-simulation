import openpyxl
X="/Users/kevindooley/Docs/School/LE/SCM and AI Oct2026 workshop/Master Data Set_4.xlsx"
wb=openpyxl.load_workbook(X,read_only=True,data_only=False)
for ws in wb.worksheets:
    print("==",ws.title, ws.max_row, ws.max_column, ws.calculate_dimension() if hasattr(ws,'calculate_dimension') else '')
