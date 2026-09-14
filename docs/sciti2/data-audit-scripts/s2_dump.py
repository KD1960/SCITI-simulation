import openpyxl
X="/Users/kevindooley/Docs/School/LE/SCM and AI Oct2026 workshop/Master Data Set_4.xlsx"
for data_only in (False,):
  wb=openpyxl.load_workbook(X,read_only=True,data_only=data_only)
  ws=wb["Glossary"]
  for r in ws.iter_rows(min_row=1,max_row=78):
    vals=[(c.coordinate,c.value) for c in r if c.value is not None]
    if vals: print(vals)
  for name,n in [("Supplier Data",8),("Shipment_CM_MFG_DC_Retailers",6),("Demand_Retailer",8)]:
    print("=====",name)
    ws=wb[name]
    for r in ws.iter_rows(min_row=1,max_row=n):
      print([(c.coordinate,c.value) for c in r if getattr(c,'value',None) is not None])
