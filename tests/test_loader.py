import json
from pathlib import Path

import numpy as np
import pandas as pd
import pytest

from sciti.data import loader

XLSX = Path.home() / "Docs/School/LE/SCM and AI Oct2026 workshop/Master Data Set_4.xlsx"


def test_parse_bom_totals_160():
    g = pd.DataFrame([[None] * 5] * 2 + [
        ["A", "Semiconductor", "SR_MCU", "x", 10], ["B", None, "SR_MOS_KIT", "x", 10],
        ["C", "Steel", "AL_Sheet", "x", 5], [None, None, "HSS_Coil", "x", 5],
        [None, "Plastic", "PP_Resin", "x", 15], [None, None, "Poly_F_EU", "x", 15],
        [None, "Glass", "LGS_4_W", "x", 10], [None, None, "TGP_5_Q", "x", 10],
        [None, "Elastomer", "EPDM", "x", 40], [None, None, "NR_201", "x", 40]])
    bom = loader.parse_bom(g)
    assert sum(v["units"] for v in bom.values()) == 160
    assert bom["TGP_5_Q"] == {"units": 10, "component": "Glass", "cm": "CM_4"}
    assert bom["NR_201"]["cm"] == "CM_3"


def test_parse_transactions_pools_missing_suppliers():
    sups = {"Supplier_1": {"sku": "SR_MCU", "cm": "CM_1", "lead_days": 10, "price": 20},
            "Supplier_2": {"sku": "SR_MCU", "cm": "CM_1", "lead_days": 12, "price": 20}}
    tx = pd.DataFrame({"Supplier": ["Supplier 1"] * 5, "Product Name": ["SR_MCU"] * 5,
                       "Lead Time": [8, 10, 12, 10, 10], "Sentiment Score": [1, 3, 3, 3, 3]})
    report = []
    loader.parse_transactions(tx, sups, report)
    assert sups["Supplier_1"]["defect_share"] == 0.2
    assert sups["Supplier_1"]["pooled"] is False
    assert sups["Supplier_2"]["pooled"] is True
    assert sups["Supplier_2"]["lead_sd_days"] == pytest.approx(np.std([8, 10, 12, 10, 10], ddof=1))
    assert any("Supplier_2" in line for line in report)


LANE_HEADER = ["Shipment_ID", "Ship_Date", "a", "b", "c", "d", "e", "Mode", "Distance (Miles)",
               "Lead Time (days)", "Quantity", "Shipping Cost/Unit", "Total_Cost", "CO2 Emissions (kg)"]


def _make_ship(rows: list) -> pd.DataFrame:
    """Repeat one lane block across all three lane-type blocks (cm_mfg/mfg_dc/dc_retail)."""
    block = pd.DataFrame([LANE_HEADER] + rows)
    gap = pd.DataFrame([[None]] * len(block))
    return pd.concat([block, gap, block, gap, block], axis=1, ignore_index=True)


def test_parse_lanes_mode_mix_and_regression():
    rows = []
    for i in range(60):
        miles = 1000 + 100 * i
        rows.append(["x", "d", "", "", "", "", "", "Air", miles, 1 + miles / 1000, 100, 10.0, 0, 0])
    for i in range(60):
        rows.append(["x", "d", "", "", "", "", "", "Ship", 5000, 20, 100, 2.0, 0, 0])
    ship = _make_ship(rows)
    lanes = loader.parse_lanes(ship, [])
    air = lanes["mfg_dc"]["modes"]["Air"]
    assert lanes["cm_mfg"]["mode_mix"] == {"Air": 0.5, "Ship": 0.5}
    assert air["lead_b"] == pytest.approx(0.001)
    assert air["lead_a"] == pytest.approx(1.0)
    assert air["cost_per_unit"] == pytest.approx(10.0)
    assert air["flat"] is False
    assert lanes["dc_retail"]["modes"]["Ship"]["lead_b"] == 0.0  # constant distance → flat


def test_fit_mode_few_rows_forces_flat():
    rows = []
    for i in range(10):
        miles = 1000 + 100 * i
        rows.append(["x", "d", "", "", "", "", "", "Air", miles, 5 + 0.01 * miles, 100, 10.0, 0, 0])
    ship = _make_ship(rows)
    report = []
    lanes = loader.parse_lanes(ship, report)
    air = lanes["mfg_dc"]["modes"]["Air"]
    lead = 5 + 0.01 * (1000 + 100 * np.arange(10))
    assert air["flat"] is True
    assert air["lead_b"] == 0.0
    assert air["lead_a"] == pytest.approx(lead.mean())
    assert any("mfg_dc/Air" in line and "flat" in line for line in report)


def test_fit_mode_negative_slope_forces_flat_even_with_enough_rows():
    rows = []
    for i in range(60):
        miles = 1000 + 100 * i
        rows.append(["x", "d", "", "", "", "", "", "Road", miles, 20 - 0.001 * miles, 100, 5.0, 0, 0])
    ship = _make_ship(rows)
    lanes = loader.parse_lanes(ship, [])
    road = lanes["mfg_dc"]["modes"]["Road"]
    assert road["flat"] is True
    assert road["lead_b"] == 0.0


def test_fit_mode_enough_rows_and_positive_slope_stays_fitted():
    rows = []
    for i in range(60):
        miles = 1000 + 100 * i
        rows.append(["x", "d", "", "", "", "", "", "Road", miles, 1 + 0.001 * miles, 100, 5.0, 0, 0])
    ship = _make_ship(rows)
    lanes = loader.parse_lanes(ship, [])
    road = lanes["mfg_dc"]["modes"]["Road"]
    assert road["flat"] is False
    assert road["lead_b"] == pytest.approx(0.001)
    assert road["lead_a"] == pytest.approx(1.0)


def test_parse_demand_shapes():
    rows = [[None] * 27] * 4
    for y in (2020, 2021):
        for w in range(1, 53):
            rows.append([y, w, "Q1"] + [float(w)] * 24)
    years, hist = loader.parse_demand(pd.DataFrame(rows))
    assert years == [2020, 2021]
    assert sorted(hist) == [f"Retail_{r}" for r in range(1, 9)]
    assert len(hist["Retail_8"]["C"]) == 104


def test_prepare_missing_sentiment_score_column(tmp_path):
    import openpyxl
    # Build minimal workbook with all sheets but missing Sentiment Score
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    # Glossary sheet (minimal)
    gs = wb.create_sheet("Glossary")
    gs.append([None] * 5)
    gs.append([None] * 5)
    for i in range(10):
        gs.append(["Component", None, f"SKU_{i}", "x", 10])

    # Supplier Data sheet (missing Sentiment Score in columns A:L)
    ss = wb.create_sheet("Supplier Data")
    # Columns A:L with all required columns EXCEPT Sentiment Score
    ss.append(["Supplier", "Lead Time", "Product Name", "x", "x", "x", "x", "x", "x", "x", "x", "x"])
    ss.append(["Supplier 1", 10, "SR_MCU", "", "", "", "", "", "", "", "", ""])

    # Shipment sheet (with proper header structure)
    header = ["Shipment_ID", "Ship_Date", "a", "b", "c", "d", "e", "Mode", "Distance (Miles)",
              "Lead Time (days)", "Quantity", "Shipping Cost/Unit", "Total_Cost", "CO2 Emissions (kg)"]
    hs = wb.create_sheet("Shipment_CM_MFG_DC_Retailers")
    hs.append(header + [None] + header + [None] + header)

    # Demand sheet
    ds = wb.create_sheet("Demand_Retailer")
    ds.append([None] * 27)
    ds.append([None] * 27)
    ds.append([None] * 27)
    ds.append([None] * 27)
    ds.append([2020, 1, "Q1"] + [100.0] * 24)

    xlsx_path = tmp_path / "missing_col.xlsx"
    wb.save(xlsx_path)

    with pytest.raises(ValueError) as exc_info:
        loader.prepare(xlsx_path, tmp_path)
    err_msg = str(exc_info.value)
    assert "Supplier Data" in err_msg
    assert "Sentiment Score" in err_msg


def test_prepare_missing_shipment_block_columns(tmp_path):
    import openpyxl
    # Build workbook where all three shipment blocks lack Distance (Miles) and Shipping Cost/Unit
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    # Glossary sheet (minimal)
    gs = wb.create_sheet("Glossary")
    gs.append([None] * 5)
    gs.append([None] * 5)
    for i in range(10):
        gs.append(["Component", None, f"SKU_{i}", "x", 10])

    # Supplier Data sheet (complete)
    ss = wb.create_sheet("Supplier Data")
    ss.append(["Supplier", "Lead Time", "Product Name", "Sentiment Score", "x", "x", "x", "x", "x", "x", "x", "x"])
    ss.append(["Supplier 1", 10, "SR_MCU", 1, "", "", "", "", "", "", "", ""])

    # Shipment sheet with blocks missing Distance (Miles) and Shipping Cost/Unit
    # Each block needs 14 columns, with gaps between blocks
    header_missing = ["Shipment_ID", "Ship_Date", "a", "b", "c", "d", "e", "Mode", "Lead Time (days)",
                      "Quantity", "Total_Cost", "CO2 Emissions (kg)", "x", "x"]  # 14 cols
    hs = wb.create_sheet("Shipment_CM_MFG_DC_Retailers")
    hs.append(header_missing + [None] + header_missing + [None] + header_missing)

    # Demand sheet
    ds = wb.create_sheet("Demand_Retailer")
    ds.append([None] * 27)
    ds.append([None] * 27)
    ds.append([None] * 27)
    ds.append([None] * 27)
    ds.append([2020, 1, "Q1"] + [100.0] * 24)

    xlsx_path = tmp_path / "missing_shipment_cols.xlsx"
    wb.save(xlsx_path)

    with pytest.raises(ValueError) as exc_info:
        loader.prepare(xlsx_path, tmp_path)
    err_msg = str(exc_info.value)
    # Verify all missing columns and all block labels are mentioned
    assert "Distance (Miles)" in err_msg
    assert "Shipping Cost/Unit" in err_msg
    assert "cm_mfg" in err_msg
    assert "mfg_dc" in err_msg
    assert "dc_retail" in err_msg


@pytest.mark.realdata
@pytest.mark.skipif(not XLSX.exists(), reason="workbook not on this machine")
def test_prepare_real_workbook(tmp_path):
    base = loader.prepare(XLSX, tmp_path)
    assert len(base["suppliers"]) == 30
    assert base["history_years"] == [2020, 2021, 2022, 2023, 2024]
    assert sum(base["demand_history"]["Retail_1"]["A"][-52:]) == pytest.approx(129239, rel=0.01)
    assert set(base["lanes"]) == {"cm_mfg", "mfg_dc", "dc_retail"}
    assert json.loads((tmp_path / "baseline.json").read_text())["source"]["xlsx_sha256"]
    assert (tmp_path / "validation_report.md").read_text().startswith("# Validation report")
